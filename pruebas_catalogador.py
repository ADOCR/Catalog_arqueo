"""Pruebas de integración locales; no requieren abrir la interfaz."""

import hashlib
import json
import os
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image

from catalogador_pozos import (
    Catalog,
    CorrectionCommittedWarning,
    App,
    FIELDS,
    FIELDS_V4,
    HEIC_ENABLED,
    LEGACY_FIELDS_V3,
    SaveCommittedWarning,
    normalized,
)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class CatalogIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=Path.cwd())
        self.root = Path(self.temp.name)
        self.source = self.root / "origen"
        self.output = self.root / "salida"
        self.source.mkdir()
        self.output.mkdir()
        self.paths = []
        for index, color in enumerate(("red", "green", "blue", "orange", "purple"), 1):
            path = self.source / f"foto_{index}.jpg"
            exif = Image.Exif()
            exif[36867] = f"2026:09:{index:02d} 10:00:00"
            Image.new("RGB", (160 + index, 100 + index), color).save(path, exif=exif, quality=92)
            self.paths.append(path)
        self.original_hashes = {path: digest(path) for path in self.paths}

    def tearDown(self):
        self.temp.cleanup()

    def assert_originals_unchanged(self):
        for path, expected in self.original_hashes.items():
            self.assertEqual(digest(path), expected)

    def save_photo(
        self, catalog, path, unit, level, view="", notes="",
        operation="2", unit_type="Pozo", subunit_type="", subunit="",
    ):
        return catalog.save(
            path, operation, unit_type, unit, level, view, notes,
            subunit_type, subunit,
        )

    def correct_photo(
        self, catalog, path, unit, level, view="", notes="",
        operation="2", unit_type="Pozo", subunit_type="", subunit="",
    ):
        return catalog.correct(
            path, operation, unit_type, unit, level, view, notes,
            subunit_type, subunit,
        )

    def test_main_workflow_resume_skip_and_failure_rollback(self):
        catalog = Catalog(self.source, self.output, "H-96-AM")
        first = self.save_photo(catalog, self.paths[0], "2705", "N3", "Perfil", "Primera")
        second = self.save_photo(catalog, self.paths[1], "2705", "N3", "Planta", "Segunda")
        self.assertEqual(first["Nombre_archivo"], "H-96-AM_OP2_P2705_N3_F001.jpg")
        self.assertEqual(second["Nombre_archivo"], "H-96-AM_OP2_P2705_N3_F002.jpg")
        self.assertEqual(digest(catalog.photos / first["Nombre_archivo"]), self.original_hashes[self.paths[0]])
        self.assertEqual(digest(catalog.photos / second["Nombre_archivo"]), self.original_hashes[self.paths[1]])

        catalog.set_skipped(self.paths[2], True)
        self.assertEqual(catalog.status_of(self.paths[2]), "skipped")
        reopened = Catalog(self.source, self.output, "H-96-AM")
        self.assertEqual(reopened.status_of(self.paths[2]), "skipped")
        reopened.set_skipped(self.paths[2], False)
        self.assertEqual(reopened.status_of(self.paths[2]), "pending")

        rows_before = len(reopened.rows)
        completed_before = set(reopened.completed)
        real_writer = reopened._write_workbook

        def blocked():
            raise PermissionError("Excel bloqueado")

        reopened._write_workbook = blocked
        with self.assertRaises(PermissionError):
            self.save_photo(reopened, self.paths[2], "2705", "N3", notes="Datos conservables")
        self.assertEqual(len(reopened.rows), rows_before)
        self.assertEqual(reopened.completed, completed_before)
        self.assertFalse((reopened.photos / "H-96-AM_OP2_P2705_N3_F003.jpg").exists())
        self.assertFalse(reopened.journal_path.exists())
        reopened._write_workbook = real_writer

        third = self.save_photo(reopened, self.paths[2], "2705", "N3")
        self.assertEqual(third["Nombre_archivo"], "H-96-AM_OP2_P2705_N3_F003.jpg")
        again = Catalog(self.source, self.output, "H-96-AM")
        self.assertEqual(len(again.rows), 3)
        self.assertEqual(len(again.completed), 3)

        workbook = load_workbook(again.xlsx_path, read_only=True, data_only=True)
        sheet = workbook.active
        self.assertEqual([cell.value for cell in sheet[1]], FIELDS)
        values = list(sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(values), 3)
        name_column = FIELDS.index("Nombre_archivo")
        for row in values:
            self.assertNotIn("\\", row[name_column])
            self.assertNotIn("/", row[name_column])
        workbook.close()
        self.assert_originals_unchanged()

    def test_auxiliary_pit_with_invalid_exif_controls_is_saved(self):
        photo = self.source / "auxiliar_exif_nulo.jpg"
        exif = Image.Exif()
        exif[36867] = "\x00" * 19
        Image.new("RGB", (120, 80), "brown").save(photo, exif=exif, quality=90)
        original_hash = digest(photo)

        catalog = Catalog(self.source, self.output, "H-96-AM")
        row = self.save_photo(
            catalog, photo, "A1", "5", unit_type="Pozo auxiliar"
        )
        self.assertEqual(row["Nombre_archivo"], "H-96-AM_OP2_PA1_5_F001.jpg")
        self.assertEqual(digest(catalog.photos / row["Nombre_archivo"]), original_hash)

        workbook = load_workbook(catalog.xlsx_path, read_only=True, data_only=True)
        saved_exif = workbook.active.cell(workbook.active.max_row, FIELDS.index("Fecha_EXIF") + 1).value
        workbook.close()
        self.assertIn(saved_exif, (None, ""))

    def test_operations_and_unit_types_use_independent_sequences(self):
        catalog = Catalog(self.source, self.output, "OTRO-SITIO")
        op2_well = self.save_photo(catalog, self.paths[0], "1", "5")
        op3_well = self.save_photo(
            catalog, self.paths[1], "1", "5", operation="3"
        )
        trench = self.save_photo(
            catalog, self.paths[2], "T1", "5", unit_type="Trinchera"
        )
        trench_second = self.save_photo(
            catalog, self.paths[3], "1", "5", unit_type="Trinchera"
        )
        auxiliary = self.save_photo(
            catalog, self.paths[4], "PA1", "5", unit_type="Pozo auxiliar"
        )

        self.assertEqual(op2_well["Nombre_archivo"], "OTRO-SITIO_OP2_P1_5_F001.jpg")
        self.assertEqual(op3_well["Nombre_archivo"], "OTRO-SITIO_OP3_P1_5_F001.jpg")
        self.assertEqual(trench["Nombre_archivo"], "OTRO-SITIO_OP2_T1_5_F001.jpg")
        self.assertEqual(trench_second["Nombre_archivo"], "OTRO-SITIO_OP2_T1_5_F002.jpg")
        self.assertEqual(auxiliary["Nombre_archivo"], "OTRO-SITIO_OP2_PA1_5_F001.jpg")
        self.assertEqual(auxiliary["Operacion"], "2")
        self.assertEqual(auxiliary["Tipo_unidad"], "Pozo auxiliar")
        self.assertEqual(auxiliary["Unidad"], "A1")

    def test_optional_subunits_have_independent_sequences(self):
        catalog = Catalog(self.source, self.output, "H-96-AM")
        square_13_first = self.save_photo(
            catalog, self.paths[0], "1", "1", operation="3", unit_type="Trinchera",
            subunit_type="Cuadro", subunit="13",
        )
        square_13_second = self.save_photo(
            catalog, self.paths[1], "1", "1", operation="3", unit_type="Trinchera",
            subunit_type="Cuadro", subunit="13",
        )
        square_14 = self.save_photo(
            catalog, self.paths[2], "1", "1", operation="3", unit_type="Trinchera",
            subunit_type="Cuadro", subunit="14",
        )
        without_subunit = self.save_photo(
            catalog, self.paths[3], "1", "1", operation="3", unit_type="Trinchera",
        )

        self.assertEqual(
            square_13_first["Nombre_archivo"], "H-96-AM_OP3_T1_CD13_1_F001.jpg"
        )
        self.assertEqual(
            square_13_second["Nombre_archivo"], "H-96-AM_OP3_T1_CD13_1_F002.jpg"
        )
        self.assertEqual(square_14["Nombre_archivo"], "H-96-AM_OP3_T1_CD14_1_F001.jpg")
        self.assertEqual(without_subunit["Nombre_archivo"], "H-96-AM_OP3_T1_1_F001.jpg")
        self.assertEqual(square_13_first["Tipo_subunidad"], "Cuadro")
        self.assertEqual(square_13_first["Subunidad"], "13")
        corrected = self.correct_photo(
            catalog, self.paths[2], "1", "1", operation="3", unit_type="Trinchera",
            subunit_type="Cuadro", subunit="13",
        )
        self.assertEqual(corrected["Nombre_archivo"], "H-96-AM_OP3_T1_CD13_1_F003.jpg")
        self.assertFalse((catalog.photos / square_14["Nombre_archivo"]).exists())
        self.assertTrue((catalog.photos / corrected["Nombre_archivo"]).is_file())
        with self.assertRaises(ValueError):
            self.save_photo(
                catalog, self.paths[4], "1", "1", unit_type="Trinchera",
                subunit_type="Cuadro", subunit="",
            )
        self.assert_originals_unchanged()

    def test_v2_2_catalog_adds_blank_optional_subunits_without_renaming(self):
        catalog = Catalog(self.source, self.output, "SITE")
        original_row = self.save_photo(catalog, self.paths[0], "1", "N1")

        workbook = load_workbook(catalog.xlsx_path)
        sheet = workbook.active
        old_values = [sheet.cell(2, index + 1).value for index in range(len(FIELDS))]
        by_name = dict(zip(FIELDS, old_values))
        sheet.delete_rows(1, sheet.max_row)
        sheet.append(FIELDS_V4)
        sheet.append([by_name[field] for field in FIELDS_V4])
        workbook.save(catalog.xlsx_path)
        workbook.close()

        reopened = Catalog(self.source, self.output, "SITE")
        migrated = reopened.rows[0]
        self.assertEqual(migrated["Tipo_subunidad"], "")
        self.assertEqual(migrated["Subunidad"], "")
        self.assertEqual(migrated["Nombre_archivo"], original_row["Nombre_archivo"])
        self.assertTrue((reopened.photos / original_row["Nombre_archivo"]).is_file())
        self.assertTrue((self.output / "respaldo_catalogo_antes_v2_3.xlsx").is_file())
        self.assertTrue((self.output / "respaldo_estado_antes_v2_3.json").is_file())
        check = load_workbook(reopened.xlsx_path, read_only=True, data_only=True)
        self.assertEqual([cell.value for cell in check.active[1]], FIELDS)
        check.close()

    def test_legacy_catalog_migrates_operation_and_unit_columns(self):
        delivery = self.output / "Entrega_Museo"
        photos = delivery / "Fotos"
        photos.mkdir(parents=True)
        old_filename = "SITE_P1_N1_F001.jpg"
        shutil.copy2(self.paths[0], photos / old_filename)

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(LEGACY_FIELDS_V3)
        sheet.append([
            "SITE", "1", "N1", 1, old_filename, self.paths[0].name,
            "2026:09:01 10:00:00", "Planta", "Registro anterior",
            "2026-09-01T10:00:00",
        ])
        workbook.save(delivery / "Catalogo_fotografico_pozos.xlsx")
        original = normalized(self.paths[0])
        (self.output / "estado_catalogador.json").write_text(json.dumps({
            "version": 3,
            "processed_paths": [original],
            "skipped_paths": [],
            "last_path": original,
            "record_map": {original: old_filename},
            "corrections": [],
        }), encoding="utf-8")

        migrated = Catalog(self.source, self.output, "SITE", legacy_operation="OP2")
        self.assertEqual(migrated.rows[0]["Operacion"], "2")
        self.assertEqual(migrated.rows[0]["Tipo_unidad"], "Pozo")
        self.assertEqual(migrated.rows[0]["Unidad"], "1")
        self.assertEqual(migrated.rows[0]["Nombre_archivo"], old_filename)
        self.assertTrue((self.output / "respaldo_catalogo_antes_v2_2.xlsx").is_file())
        self.assertTrue((self.output / "respaldo_estado_antes_v2_2.json").is_file())
        reopened = Catalog(self.source, self.output, "SITE")
        self.assertEqual(reopened.rows, migrated.rows)

        check = load_workbook(migrated.xlsx_path, read_only=True, data_only=True)
        self.assertEqual([cell.value for cell in check.active[1]], FIELDS)
        check.close()

    def test_old_state_compatibility(self):
        catalog = Catalog(self.source, self.output, "SITE")
        self.save_photo(catalog, self.paths[0], "1", "N1")
        old_state = {"processed_paths": sorted(catalog.completed)}
        catalog.state_path.write_text(json.dumps(old_state), encoding="utf-8")
        reopened = Catalog(self.source, self.output, "SITE")
        self.assertEqual(reopened.completed, {normalized(self.paths[0])})
        self.assertEqual(reopened.skipped, set())

    def test_interrupted_state_is_recovered_from_journal(self):
        catalog = Catalog(self.source, self.output, "SITE")
        real_state_writer = catalog._write_state

        def state_blocked(*_args, **_kwargs):
            raise PermissionError("Estado bloqueado")

        catalog._write_state = state_blocked
        with self.assertRaises(SaveCommittedWarning):
            self.save_photo(catalog, self.paths[0], "2", "N4")
        self.assertTrue(catalog.journal_path.exists())
        self.assertTrue((catalog.photos / "SITE_OP2_P2_N4_F001.jpg").exists())
        catalog._write_state = real_state_writer

        recovered = Catalog(self.source, self.output, "SITE")
        self.assertFalse(recovered.journal_path.exists())
        self.assertEqual(len(recovered.rows), 1)
        self.assertEqual(recovered.completed, {normalized(self.paths[0])})
        self.assert_originals_unchanged()

    def test_correction_updates_row_copy_and_history_without_touching_original(self):
        catalog = Catalog(self.source, self.output, "SITE")
        first = self.save_photo(catalog, self.paths[0], "1", "N1", "Planta", "Dato inicial")
        self.save_photo(catalog, self.paths[1], "2", "N2", "Perfil", "Otra foto")

        same_group = self.correct_photo(
            catalog, self.paths[0], "1", "N1", "Detalle", "Observación corregida"
        )
        self.assertEqual(same_group["Nombre_archivo"], first["Nombre_archivo"])
        self.assertEqual(same_group["Foto_en_unidad_y_nivel"], "1")
        self.assertEqual(same_group["Vista_o_detalle"], "Detalle")
        self.assertEqual(same_group["Observaciones"], "Observación corregida")

        moved = self.correct_photo(
            catalog, self.paths[0], "2", "N2", "Detalle", "Pozo corregido"
        )
        self.assertEqual(moved["Nombre_archivo"], "SITE_OP2_P2_N2_F002.jpg")
        self.assertFalse((catalog.photos / first["Nombre_archivo"]).exists())
        self.assertTrue((catalog.photos / moved["Nombre_archivo"]).is_file())
        self.assertEqual(
            digest(catalog.photos / moved["Nombre_archivo"]), self.original_hashes[self.paths[0]]
        )
        self.assertEqual(len(catalog.rows), 2)
        self.assertEqual(len(catalog.corrections), 2)
        self.assertEqual(catalog.corrections[-1]["before"]["Unidad"], "1")
        self.assertEqual(catalog.corrections[-1]["after"]["Unidad"], "2")

        reopened = Catalog(self.source, self.output, "SITE")
        located = reopened.row_for_path(self.paths[0])
        self.assertIsNotNone(located)
        self.assertEqual(located[1]["Nombre_archivo"], "SITE_OP2_P2_N2_F002.jpg")
        self.assertEqual(len(reopened.corrections), 2)
        self.assert_originals_unchanged()

    def test_correction_rolls_back_if_excel_cannot_be_written(self):
        catalog = Catalog(self.source, self.output, "SITE")
        original_row = self.save_photo(catalog, self.paths[0], "1", "N1", "Planta", "Original")
        real_writer = catalog._write_workbook

        def blocked():
            raise PermissionError("Excel bloqueado")

        catalog._write_workbook = blocked
        with self.assertRaises(PermissionError):
            self.correct_photo(
                catalog, self.paths[0], "3", "N4", "Perfil", "No debe quedar"
            )
        catalog._write_workbook = real_writer

        self.assertTrue((catalog.photos / original_row["Nombre_archivo"]).is_file())
        self.assertFalse((catalog.photos / "SITE_OP2_P3_N4_F001.jpg").exists())
        self.assertEqual(catalog.rows[0], original_row)
        self.assertEqual(catalog.record_map[normalized(self.paths[0])], original_row["Nombre_archivo"])
        self.assertEqual(catalog.corrections, [])
        self.assertFalse(catalog.journal_path.exists())

        reopened = Catalog(self.source, self.output, "SITE")
        self.assertEqual(reopened.rows[0], original_row)
        self.assert_originals_unchanged()

    def test_interrupted_correction_state_is_recovered_from_journal(self):
        catalog = Catalog(self.source, self.output, "SITE")
        self.save_photo(catalog, self.paths[0], "1", "N1", "Planta", "Inicial")
        real_state_writer = catalog._write_state

        def state_blocked(*_args, **_kwargs):
            raise PermissionError("Estado bloqueado")

        catalog._write_state = state_blocked
        with self.assertRaises(CorrectionCommittedWarning):
            self.correct_photo(
                catalog, self.paths[0], "4", "N8", "Perfil", "=Corregida"
            )
        self.assertTrue(catalog.journal_path.exists())
        self.assertTrue((catalog.photos / "SITE_OP2_P4_N8_F001.jpg").is_file())
        catalog._write_state = real_state_writer

        recovered = Catalog(self.source, self.output, "SITE")
        self.assertFalse(recovered.journal_path.exists())
        located = recovered.row_for_path(self.paths[0])
        self.assertIsNotNone(located)
        self.assertEqual(located[1]["Nombre_archivo"], "SITE_OP2_P4_N8_F001.jpg")
        self.assertEqual(located[1]["Observaciones"], "'=Corregida")
        self.assertEqual(len(recovered.corrections), 1)
        self.assert_originals_unchanged()

    def test_old_state_rebuilds_unique_record_mapping_for_corrections(self):
        catalog = Catalog(self.source, self.output, "SITE")
        row = self.save_photo(catalog, self.paths[0], "1", "N1")
        old_state = {
            "version": 2,
            "processed_paths": sorted(catalog.completed),
            "skipped_paths": [],
            "last_path": "",
        }
        catalog.state_path.write_text(json.dumps(old_state), encoding="utf-8")

        reopened = Catalog(self.source, self.output, "SITE")
        located = reopened.row_for_path(self.paths[0])
        self.assertIsNotNone(located)
        self.assertEqual(located[1]["Nombre_archivo"], row["Nombre_archivo"])
        corrected = self.correct_photo(
            reopened, self.paths[0], "1", "N1", "Detalle", "Recuperada"
        )
        self.assertEqual(corrected["Nombre_archivo"], row["Nombre_archivo"])

    @unittest.skipUnless(os.name == "nt", "prueba visual de widgets disponible en Windows")
    def test_correction_mode_ui_smoke(self):
        catalog = Catalog(self.source, self.output, "SITE")
        self.save_photo(
            catalog, self.paths[0], "1", "N2", "Perfil", "Revisar",
            unit_type="Trinchera", subunit_type="Cuadro", subunit="13",
        )
        app = App()
        app.withdraw()
        try:
            app.catalog = catalog
            catalog.set_skipped(self.paths[1], True)
            app._update_progress()
            self.assertEqual(
                app.progress_var.get(),
                "Registradas: 1 · Pendientes: 3 · Omitidas: 1 · Falta revisar: 80,0 %",
            )
            app._select_photo(self.paths[0], remember=False)
            app._begin_correction()
            app.update_idletasks()
            self.assertEqual(app.editing_path, self.paths[0])
            self.assertEqual(app.operation_var.get(), "2")
            self.assertEqual(app.unit_type_var.get(), "Trinchera")
            self.assertEqual(app.pit_var.get(), "T1")
            self.assertEqual(app.subunit_type_var.get(), "Cuadro")
            self.assertEqual(app.subunit_var.get(), "13")
            self.assertEqual(app.level_var.get(), "N2")
            self.assertEqual(app.save_button.cget("text"), "Guardar corrección")
            self.assertEqual(app.skip_button.cget("text"), "Cancelar corrección")
            self.assertEqual(app.proposed_var.get(), "SITE_OP2_T1_CD13_N2_F001.jpg")
        finally:
            app._close()

    @unittest.skipUnless(HEIC_ENABLED, "pillow-heif no está instalado en este entorno")
    def test_heic_copy_and_catalog(self):
        heic_path = self.source / "foto_heic.heic"
        Image.new("RGB", (96, 64), "teal").save(heic_path, format="HEIF", quality=90)
        original_hash = digest(heic_path)
        catalog = Catalog(self.source, self.output, "SITE")
        self.assertIn(heic_path, catalog.files)
        row = self.save_photo(catalog, heic_path, "7", "N2")
        copy = catalog.photos / row["Nombre_archivo"]
        self.assertEqual(copy.suffix, ".heic")
        self.assertEqual(digest(copy), original_hash)
        with Image.open(copy) as image:
            image.load()
            self.assertEqual(image.size, (96, 64))

    @unittest.skipUnless(os.name == "nt", "prueba específica del bloqueo de archivos de Windows")
    def test_real_windows_excel_lock_preserves_pending_photo(self):
        import ctypes

        catalog = Catalog(self.source, self.output, "SITE")
        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        create_file = kernel32.CreateFileW
        create_file.argtypes = [
            ctypes.c_wchar_p, ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p,
            ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p,
        ]
        create_file.restype = ctypes.c_void_p
        close_handle = kernel32.CloseHandle
        close_handle.argtypes = [ctypes.c_void_p]
        close_handle.restype = ctypes.c_int
        handle = create_file(
            str(catalog.xlsx_path),
            0x80000000,  # GENERIC_READ
            0,           # sin uso compartido, como un bloqueo exclusivo
            None,
            3,           # OPEN_EXISTING
            0x80,        # FILE_ATTRIBUTE_NORMAL
            None,
        )
        self.assertNotIn(handle, (None, ctypes.c_void_p(-1).value))
        try:
            with self.assertRaises(OSError):
                self.save_photo(
                    catalog, self.paths[0], "9", "N5", notes="Debe permanecer pendiente"
                )
        finally:
            close_handle(handle)

        self.assertEqual(catalog.rows, [])
        self.assertEqual(catalog.completed, set())
        self.assertFalse((catalog.photos / "SITE_OP2_P9_N5_F001.jpg").exists())
        self.assertFalse(catalog.journal_path.exists())
        row = self.save_photo(catalog, self.paths[0], "9", "N5", notes="Reintento")
        self.assertEqual(row["Nombre_archivo"], "SITE_OP2_P9_N5_F001.jpg")

        handle = create_file(
            str(catalog.xlsx_path), 0x80000000, 0, None, 3, 0x80, None
        )
        self.assertNotIn(handle, (None, ctypes.c_void_p(-1).value))
        try:
            with self.assertRaises(OSError):
                self.correct_photo(
                    catalog, self.paths[0], "10", "N6", "Perfil", "No debe aplicarse"
                )
        finally:
            close_handle(handle)

        self.assertTrue((catalog.photos / row["Nombre_archivo"]).is_file())
        self.assertFalse((catalog.photos / "SITE_OP2_P10_N6_F001.jpg").exists())
        self.assertEqual(catalog.rows[0], row)
        corrected = self.correct_photo(
            catalog, self.paths[0], "10", "N6", "Perfil", "Reintento"
        )
        self.assertEqual(corrected["Nombre_archivo"], "SITE_OP2_P10_N6_F001.jpg")


if __name__ == "__main__":
    unittest.main(verbosity=2)
