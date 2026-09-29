"""Catalogador local de fotografías arqueológicas para Windows.

Los originales nunca se modifican. La rotación, el zoom y el desplazamiento
pertenecen únicamente al visor.
"""

from __future__ import annotations

import csv
import json
import math
import os
import re
import shutil
import tempfile
import tkinter as tk
from collections import OrderedDict, deque
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from typing import Callable

from PIL import Image, ImageDraw, ImageOps, ImageTk
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from pillow_heif import register_heif_opener

    register_heif_opener()
    HEIC_ENABLED = True
except ImportError:
    HEIC_ENABLED = False


APP_DIR = Path(__file__).resolve().parent
ASSET_PATH = APP_DIR / "assets" / "cucharilla.png"
STATE_VERSION = 3
SUFFIXES = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".webp"}
if HEIC_ENABLED:
    SUFFIXES.update({".heic", ".heif"})

FIELDS = [
    "Sitio", "Pozo", "Nivel", "Foto_en_pozo_y_nivel", "Nombre_archivo",
    "Nombre_original", "Fecha_EXIF", "Vista_o_detalle", "Observaciones",
    "Fecha_catalogacion",
]
OLD_FIELDS = [
    "Sitio", "Pozo", "Nivel", "Foto_en_pozo_y_nivel", "Nombre_nuevo",
    "Ruta_relativa", "Nombre_original", "Ruta_original", "Fecha_EXIF",
    "Vista_o_detalle", "Observaciones", "Fecha_catalogacion",
]

SAND = "#F5F2EB"
WHITE = "#FFFFFF"
INK = "#24312D"
MUTED = "#65716C"
GREEN = "#2F5D50"
GREEN_DARK = "#24493F"
GREEN_PALE = "#E8F0ED"
GOLD = "#B8833B"
RED = "#A94A45"
CANVAS_BG = "#202624"
ILLEGAL_EXCEL_CHARACTERS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


def enable_windows_dpi_awareness() -> None:
    """Permite que Tk reciba dimensiones correctas con escalado de Windows."""
    if os.name != "nt":
        return
    try:
        import ctypes

        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except Exception:
            pass


def excel_safe(value: object) -> str:
    """Convierte a texto válido para XLSX y evita fórmulas accidentales."""
    # Algunas cámaras escriben fechas EXIF rellenadas con NUL. Esos controles
    # no son válidos en XML/XLSX, aunque el archivo de imagen sí pueda abrirse.
    text = ILLEGAL_EXCEL_CHARACTERS.sub("", str(value or ""))
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def normalized(path: Path) -> str:
    return str(path.resolve())


def discover(source: Path, output: Path) -> list[Path]:
    files: list[Path] = []
    output_resolved = output.resolve()
    for path in source.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in SUFFIXES:
            continue
        try:
            if path.resolve().is_relative_to(output_resolved):
                continue
        except (OSError, ValueError):
            continue
        files.append(path)
    return sorted(files, key=lambda item: (str(item.parent).casefold(), item.name.casefold()))


class SaveCommittedWarning(Exception):
    """La foto y el Excel se guardaron, pero el estado requiere recuperación."""


class CorrectionCommittedWarning(Exception):
    """La corrección se aplicó, pero el estado requiere recuperación."""


class Catalog:
    """Persistencia del catálogo y transacciones de copia/Excel/estado."""

    def __init__(self, source: Path, output: Path, site: str):
        self.source = source.resolve()
        self.output = output.resolve()
        self.site = site
        self.delivery = self.output / "Entrega_Museo"
        self.photos = self.delivery / "Fotos"
        self.xlsx_path = self.delivery / "Catalogo_fotografico_pozos.xlsx"
        self.state_path = self.output / "estado_catalogador.json"
        self.journal_path = self.output / ".transaccion_catalogador.json"
        self.output.mkdir(parents=True, exist_ok=True)
        self.photos.mkdir(parents=True, exist_ok=True)
        self.rows: list[dict[str, str]] = []
        self.completed: set[str] = set()
        self.skipped: set[str] = set()
        self.record_map: dict[str, str] = {}
        self.corrections: list[dict[str, object]] = []
        self.last_path = ""

        if not self.xlsx_path.exists() and not self.state_path.exists():
            self._migrate_previous()
        self._open_workbook()
        self._load_state()
        self._recover_interrupted_transaction()
        self.files = discover(self.source, self.output)
        self._rebuild_record_map()

    def _open_workbook(self) -> None:
        if self.xlsx_path.exists():
            self.wb = load_workbook(self.xlsx_path)
            self.ws = self.wb.active
            if [cell.value for cell in self.ws[1]] != FIELDS:
                raise ValueError("El Excel existente tiene columnas diferentes.")
            self.rows = [
                dict(zip(FIELDS, [str(value or "") for value in values]))
                for values in self.ws.iter_rows(min_row=2, values_only=True)
                if any(value is not None for value in values)
            ]
            return

        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "Fotografías"
        self.ws.append(FIELDS)
        self.ws.freeze_panes = "D2"
        self.ws.sheet_view.showGridLines = False
        self.ws.row_dimensions[1].height = 31
        for cell in self.ws[1]:
            cell.fill = PatternFill("solid", fgColor="2F5D50")
            cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", horizontal="center")
        widths = [14, 12, 12, 23, 43, 35, 23, 34, 52, 22]
        for index, width in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(index)].width = width
        self._write_workbook()

    def _load_state(self) -> None:
        if self.state_path.exists():
            try:
                state = json.loads(self.state_path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError) as exc:
                raise ValueError(f"No se pudo leer el estado del catálogo: {exc}") from exc
            # Compatibilidad: la versión anterior solo tenía processed_paths.
            self.completed = set(state.get("processed_paths", []))
            self.skipped = set(state.get("skipped_paths", [])) - self.completed
            self.last_path = str(state.get("last_path", ""))
            self.record_map = {
                str(path): str(filename)
                for path, filename in dict(state.get("record_map", {})).items()
            }
            self.corrections = list(state.get("corrections", []))
            if len(self.completed) != len(self.rows) and not self.journal_path.exists():
                raise ValueError(
                    "El estado y el Excel no coinciden. Haz una copia de la carpeta "
                    "de trabajo antes de corregirlos."
                )
        elif self.rows and not self.journal_path.exists():
            raise ValueError(
                "Falta estado_catalogador.json; se necesita para reanudar sin repetir fotos."
            )

    def _state_data(
        self,
        completed: set[str] | None = None,
        skipped: set[str] | None = None,
        last_path: str | None = None,
        record_map: dict[str, str] | None = None,
        corrections: list[dict[str, object]] | None = None,
    ) -> dict[str, object]:
        return {
            "version": STATE_VERSION,
            "processed_paths": sorted(self.completed if completed is None else completed),
            "skipped_paths": sorted(self.skipped if skipped is None else skipped),
            "last_path": self.last_path if last_path is None else last_path,
            "record_map": dict(sorted(
                (self.record_map if record_map is None else record_map).items()
            )),
            "corrections": self.corrections if corrections is None else corrections,
        }

    def _atomic_json(self, path: Path, data: dict[str, object]) -> None:
        fd, temp_name = tempfile.mkstemp(prefix=path.stem + "_", suffix=".tmp", dir=path.parent)
        os.close(fd)
        temp_path = Path(temp_name)
        try:
            temp_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
            os.replace(temp_path, path)
        finally:
            temp_path.unlink(missing_ok=True)

    def _write_state(
        self,
        completed: set[str] | None = None,
        skipped: set[str] | None = None,
        last_path: str | None = None,
        record_map: dict[str, str] | None = None,
        corrections: list[dict[str, object]] | None = None,
    ) -> None:
        self._atomic_json(
            self.state_path,
            self._state_data(completed, skipped, last_path, record_map, corrections),
        )

    def remember_current(self, path: Path | None) -> None:
        value = normalized(path) if path else ""
        self._write_state(last_path=value)
        self.last_path = value

    def _write_journal(self, data: dict[str, object]) -> None:
        self._atomic_json(self.journal_path, data)

    def _recover_interrupted_transaction(self) -> None:
        """Finaliza o revierte una operación cortada entre Excel y estado."""
        if not self.journal_path.exists():
            return
        try:
            journal = json.loads(self.journal_path.read_text(encoding="utf-8"))
            original = str(journal["original"])
            operation = str(journal.get("operation", "save"))
            filename = str(journal.get("filename", journal.get("new_filename", "")))
        except Exception as exc:
            raise ValueError(
                "Existe una transacción incompleta que no se puede interpretar: "
                f"{self.journal_path.name}. {exc}"
            ) from exc

        if operation == "correction":
            self._recover_interrupted_correction(journal)
            return

        destination = self.photos / filename
        excel_has_row = any(row.get("Nombre_archivo") == filename for row in self.rows)
        if excel_has_row and destination.is_file():
            completed = self.completed | {original}
            skipped = self.skipped - {original}
            record_map = dict(self.record_map)
            record_map[original] = filename
            self._write_state(
                completed=completed, skipped=skipped, last_path=original,
                record_map=record_map,
            )
            self.completed, self.skipped, self.last_path = completed, skipped, original
            self.record_map = record_map
            self.journal_path.unlink(missing_ok=True)
            return

        if destination.exists() and not excel_has_row:
            destination.unlink()
        self.journal_path.unlink(missing_ok=True)

    def _recover_interrupted_correction(self, journal: dict[str, object]) -> None:
        original = str(journal["original"])
        old_filename = str(journal["old_filename"])
        new_filename = str(journal["new_filename"])
        new_row = dict(journal["new_row"])
        history = dict(journal["history"])
        old_destination = self.photos / old_filename
        new_destination = self.photos / new_filename
        # Comparar toda la fila evita dar por aplicada una corrección de
        # observaciones cuando el proceso se interrumpió antes de escribir Excel.
        def matches_written_row(row: dict[str, str]) -> bool:
            for key in FIELDS:
                expected = str(new_row.get(key, ""))
                if key != "Foto_en_pozo_y_nivel":
                    expected = excel_safe(expected)
                if str(row.get(key, "")) != expected:
                    return False
            return True

        excel_has_new = any(matches_written_row(row) for row in self.rows)

        if excel_has_new and new_destination.is_file():
            record_map = dict(self.record_map)
            record_map[original] = new_filename
            corrections = list(self.corrections)
            if history not in corrections:
                corrections.append(history)
            self._write_state(
                last_path=original, record_map=record_map, corrections=corrections
            )
            self.record_map, self.corrections, self.last_path = record_map, corrections, original
            self.journal_path.unlink(missing_ok=True)
            return

        # Si Excel no confirmó el cambio, restaura el nombre anterior de la copia.
        if new_filename != old_filename and new_destination.exists() and not old_destination.exists():
            os.replace(new_destination, old_destination)
        self.journal_path.unlink(missing_ok=True)

    def _migrate_previous(self) -> None:
        """Importa una vez la salida anterior, sin modificarla."""
        previous_xlsx = self.output / "Catalogo_fotografico_pozos.xlsx"
        previous_csv = self.output / "Catalogo_fotografico_pozos.csv"
        old_rows: list[dict[str, str]] = []
        if previous_xlsx.exists():
            old = load_workbook(previous_xlsx, read_only=True, data_only=True).active
            if [cell.value for cell in old[1]] != OLD_FIELDS:
                raise ValueError("El Excel anterior tiene columnas diferentes; revisa la migración.")
            old_rows = [
                dict(zip(OLD_FIELDS, [str(value or "") for value in values]))
                for values in old.iter_rows(min_row=2, values_only=True)
                if any(value is not None for value in values)
            ]
        elif previous_csv.exists():
            with previous_csv.open("r", encoding="utf-8-sig", newline="") as file:
                reader = csv.DictReader(file)
                if reader.fieldnames != OLD_FIELDS:
                    raise ValueError("El CSV anterior tiene columnas diferentes; revisa la migración.")
                old_rows = list(reader)
        if not old_rows:
            return

        for row in old_rows:
            previous_copy = self.output / row["Ruta_relativa"]
            if not previous_copy.is_file():
                raise FileNotFoundError(f"Falta la copia previa: {row['Nombre_nuevo']}")
            destination = self.photos / row["Nombre_nuevo"]
            if not destination.exists():
                shutil.copy2(previous_copy, destination)

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(FIELDS)
        for row in old_rows:
            sheet.append([
                row.get("Nombre_nuevo", "") if key == "Nombre_archivo" else row.get(key, "")
                for key in FIELDS
            ])
        workbook.save(self.xlsx_path)
        completed = {row["Ruta_original"] for row in old_rows}
        record_map = {row["Ruta_original"]: row["Nombre_nuevo"] for row in old_rows}
        self._atomic_json(self.state_path, {
            "version": STATE_VERSION,
            "processed_paths": sorted(completed),
            "skipped_paths": [],
            "last_path": "",
            "record_map": record_map,
            "corrections": [],
        })

    def _append_row(self, row: dict[str, str]) -> None:
        self.ws.append([
            int(row[key]) if key == "Foto_en_pozo_y_nivel" else excel_safe(row[key])
            for key in FIELDS
        ])
        line = self.ws.max_row
        self.ws.cell(line, FIELDS.index("Foto_en_pozo_y_nivel") + 1).alignment = Alignment(horizontal="right")

    def _write_workbook(self) -> None:
        self.ws.auto_filter.ref = f"A1:J{self.ws.max_row}"
        fd, temp_name = tempfile.mkstemp(prefix="catalogo_", suffix=".xlsx", dir=self.delivery)
        os.close(fd)
        temp_path = Path(temp_name)
        try:
            self.wb.save(temp_path)
            os.replace(temp_path, self.xlsx_path)
        finally:
            temp_path.unlink(missing_ok=True)

    @staticmethod
    def validate_identifier(value: str, label: str) -> str:
        clean = value.strip()
        if not re.fullmatch(r"[A-Za-z0-9_-]{1,30}", clean):
            raise ValueError(f"{label} solo admite letras, números, guion y guion bajo (máximo 30).")
        return clean

    def remaining(self) -> list[Path]:
        return [path for path in self.files
                if normalized(path) not in self.completed and normalized(path) not in self.skipped]

    def omitted(self) -> list[Path]:
        return [path for path in self.files if normalized(path) in self.skipped]

    def status_of(self, path: Path) -> str:
        value = normalized(path)
        if value in self.completed:
            return "registered"
        if value in self.skipped:
            return "skipped"
        return "pending"

    def _rebuild_record_map(self) -> None:
        """Completa asociaciones inequívocas de estados creados por versiones anteriores."""
        filenames = {row["Nombre_archivo"] for row in self.rows}
        self.record_map = {
            path: filename for path, filename in self.record_map.items()
            if path in self.completed and filename in filenames
        }
        rows_by_original: dict[str, list[dict[str, str]]] = {}
        paths_by_name: dict[str, list[str]] = {}
        for row in self.rows:
            rows_by_original.setdefault(row["Nombre_original"], []).append(row)
        for path in self.completed:
            paths_by_name.setdefault(Path(path).name, []).append(path)
        for name, paths in paths_by_name.items():
            candidates = rows_by_original.get(name, [])
            if len(paths) == 1 and len(candidates) == 1:
                self.record_map.setdefault(paths[0], candidates[0]["Nombre_archivo"])

    def row_for_path(self, path: Path) -> tuple[int, dict[str, str]] | None:
        original = normalized(path)
        filename = self.record_map.get(original)
        if filename:
            for index, row in enumerate(self.rows):
                if row["Nombre_archivo"] == filename:
                    return index, row
        matches = [
            (index, row) for index, row in enumerate(self.rows)
            if row["Nombre_original"] == path.name
        ]
        completed_same_name = [value for value in self.completed if Path(value).name == path.name]
        if len(matches) == 1 and len(completed_same_name) == 1:
            self.record_map[original] = matches[0][1]["Nombre_archivo"]
            return matches[0]
        return None

    def next_number(self, pit: str, level: str) -> int:
        numbers = [int(row["Foto_en_pozo_y_nivel"]) for row in self.rows
                   if row["Pozo"] == pit and row["Nivel"] == level
                   and str(row["Foto_en_pozo_y_nivel"]).isdigit()]
        return max(numbers, default=0) + 1

    def proposed_name(self, path: Path, pit: str, level: str) -> str:
        pit = self.validate_identifier(pit, "Pozo")
        level = self.validate_identifier(level, "Nivel")
        safe_site = self.validate_identifier(self.site, "Código del sitio")
        number = self.next_number(pit, level)
        return f"{safe_site}_P{pit}_{level}_F{number:03d}{path.suffix.lower()}"

    def proposed_correction_name(self, path: Path, pit: str, level: str) -> str:
        pit = self.validate_identifier(pit, "Pozo")
        level = self.validate_identifier(level, "Nivel")
        located = self.row_for_path(path)
        if located is None:
            raise ValueError(
                "No se pudo asociar esta fotografía con una única fila del Excel. "
                "Puede ocurrir con catálogos antiguos que contienen nombres originales repetidos."
            )
        _, old_row = located
        if old_row["Pozo"] == pit and old_row["Nivel"] == level:
            number = int(old_row["Foto_en_pozo_y_nivel"])
        else:
            number = self.next_number(pit, level)
        safe_site = self.validate_identifier(self.site, "Código del sitio")
        return f"{safe_site}_P{pit}_{level}_F{number:03d}{path.suffix.lower()}"

    def _replace_worksheet_row(self, row_index: int, row: dict[str, str]) -> None:
        excel_row = row_index + 2
        for column, key in enumerate(FIELDS, 1):
            value: object = int(row[key]) if key == "Foto_en_pozo_y_nivel" else excel_safe(row[key])
            self.ws.cell(excel_row, column, value)
        self.ws.cell(excel_row, FIELDS.index("Foto_en_pozo_y_nivel") + 1).alignment = Alignment(
            horizontal="right"
        )

    def set_skipped(self, path: Path, skipped: bool) -> None:
        original = normalized(path)
        if original in self.completed:
            raise ValueError("La fotografía ya está registrada y no se puede omitir.")
        new_skipped = set(self.skipped)
        if skipped:
            new_skipped.add(original)
        else:
            new_skipped.discard(original)
        self._write_state(skipped=new_skipped, last_path=original)
        self.skipped = new_skipped
        self.last_path = original

    def save(self, path: Path, pit: str, level: str, view: str, notes: str) -> dict[str, str]:
        pit = self.validate_identifier(pit, "Pozo")
        level = self.validate_identifier(level, "Nivel")
        original = normalized(path)
        if original in self.completed:
            raise ValueError("Esta fotografía ya está registrada.")

        number = self.next_number(pit, level)
        filename = self.proposed_name(path, pit, level)
        destination = self.photos / filename
        if destination.exists():
            raise FileExistsError(f"Ya existe {filename}. Revisa la carpeta antes de continuar.")

        exif_date = ""
        try:
            with Image.open(path) as image:
                exif = image.getexif()
                exif_date = str(exif.get(36867) or exif.get(306) or "")
        except Exception:
            pass

        row = dict(zip(FIELDS, [
            self.site, pit, level, str(number), filename, path.name, exif_date,
            view.strip(), notes.strip(), datetime.now().isoformat(timespec="seconds"),
        ]))
        self._write_journal({
            "version": 1, "operation": "save", "original": original, "filename": filename,
            "started_at": datetime.now().isoformat(timespec="seconds"),
        })
        copied = False
        row_appended = False
        workbook_written = False
        try:
            shutil.copy2(path, destination)
            copied = True
            self._append_row(row)
            row_appended = True
            self._write_workbook()
            workbook_written = True

            completed = self.completed | {original}
            skipped = self.skipped - {original}
            record_map = dict(self.record_map)
            record_map[original] = filename
            try:
                self._write_state(
                    completed=completed, skipped=skipped, last_path=original,
                    record_map=record_map,
                )
            except Exception as exc:
                self.rows.append(row)
                self.completed, self.skipped, self.last_path = completed, skipped, original
                self.record_map = record_map
                raise SaveCommittedWarning(
                    "La copia y la fila de Excel se guardaron. No se pudo actualizar el estado; "
                    "no repitas la foto. Al reabrir, el programa intentará recuperarlo."
                ) from exc

            self.rows.append(row)
            self.completed, self.skipped, self.last_path = completed, skipped, original
            self.record_map = record_map
            self.journal_path.unlink(missing_ok=True)
            return row
        except SaveCommittedWarning:
            raise
        except Exception:
            if row_appended and not workbook_written and self.ws.max_row > len(self.rows) + 1:
                self.ws.delete_rows(self.ws.max_row)
            if copied and not workbook_written:
                destination.unlink(missing_ok=True)
            if not workbook_written:
                self.journal_path.unlink(missing_ok=True)
            raise

    def correct(
        self, path: Path, pit: str, level: str, view: str, notes: str
    ) -> dict[str, str]:
        """Corrige una fila y el nombre de su copia sin modificar el original."""
        pit = self.validate_identifier(pit, "Pozo")
        level = self.validate_identifier(level, "Nivel")
        original = normalized(path)
        if original not in self.completed:
            raise ValueError("Solo se pueden corregir fotografías ya registradas.")
        located = self.row_for_path(path)
        if located is None:
            raise ValueError(
                "No se pudo identificar de forma inequívoca la fila de esta fotografía."
            )
        row_index, old_row = located
        old_row = dict(old_row)
        old_filename = old_row["Nombre_archivo"]
        new_filename = self.proposed_correction_name(path, pit, level)
        old_destination = self.photos / old_filename
        new_destination = self.photos / new_filename
        if not old_destination.is_file():
            raise FileNotFoundError(f"Falta la copia registrada: {old_filename}")
        if new_filename != old_filename and new_destination.exists():
            raise FileExistsError(f"Ya existe {new_filename}; no se aplicó la corrección.")

        if old_row["Pozo"] == pit and old_row["Nivel"] == level:
            number = old_row["Foto_en_pozo_y_nivel"]
        else:
            number = str(self.next_number(pit, level))
        new_row = dict(old_row)
        new_row.update({
            "Pozo": pit,
            "Nivel": level,
            "Foto_en_pozo_y_nivel": str(number),
            "Nombre_archivo": new_filename,
            "Vista_o_detalle": view.strip(),
            "Observaciones": notes.strip(),
        })
        history: dict[str, object] = {
            "original_path": original,
            "corrected_at": datetime.now().isoformat(timespec="seconds"),
            "before": old_row,
            "after": new_row,
        }
        self._write_journal({
            "version": 1,
            "operation": "correction",
            "original": original,
            "old_filename": old_filename,
            "new_filename": new_filename,
            "old_row": old_row,
            "new_row": new_row,
            "history": history,
        })

        renamed = False
        workbook_written = False
        try:
            if new_filename != old_filename:
                os.replace(old_destination, new_destination)
                renamed = True
            self._replace_worksheet_row(row_index, new_row)
            self._write_workbook()
            workbook_written = True
            record_map = dict(self.record_map)
            record_map[original] = new_filename
            corrections = list(self.corrections) + [history]
            try:
                self._write_state(
                    last_path=original, record_map=record_map, corrections=corrections
                )
            except Exception as exc:
                self.rows[row_index] = new_row
                self.record_map, self.corrections, self.last_path = record_map, corrections, original
                raise CorrectionCommittedWarning(
                    "La corrección quedó aplicada en la copia y el Excel. El estado se "
                    "completará automáticamente al reabrir."
                ) from exc
            self.rows[row_index] = new_row
            self.record_map, self.corrections, self.last_path = record_map, corrections, original
            self.journal_path.unlink(missing_ok=True)
            return new_row
        except CorrectionCommittedWarning:
            raise
        except Exception:
            if not workbook_written:
                self._replace_worksheet_row(row_index, old_row)
                if renamed and new_destination.exists() and not old_destination.exists():
                    os.replace(new_destination, old_destination)
                self.journal_path.unlink(missing_ok=True)
            raise


class ImageViewer(ttk.Frame):
    """Visor de alta resolución que solo reescala el recorte visible."""

    MIN_ZOOM = 0.001
    MAX_ZOOM = 8.0

    def __init__(self, master: tk.Misc, zoom_callback: Callable[[int], None]):
        super().__init__(master, style="Panel.TFrame")
        self.zoom_callback = zoom_callback
        self.canvas = tk.Canvas(self, background=CANVAS_BG, highlightthickness=0, cursor="fleur")
        self.canvas.pack(fill="both", expand=True)
        self._base: Image.Image | None = None
        self._image: Image.Image | None = None
        self._photo: ImageTk.PhotoImage | None = None
        self._path: Path | None = None
        self._rotation = 0
        self._zoom = 1.0
        self._pan_x = 0.0
        self._pan_y = 0.0
        self._drag_start: tuple[int, int, float, float] | None = None
        self._resize_job: str | None = None

        self.canvas.bind("<Configure>", self._on_configure)
        self.canvas.bind("<MouseWheel>", self._on_wheel)
        self.canvas.bind("<Control-MouseWheel>", self._on_wheel)
        self.canvas.bind("<Button-4>", lambda event: self._on_linux_wheel(event, 1))
        self.canvas.bind("<Button-5>", lambda event: self._on_linux_wheel(event, -1))
        self.canvas.bind("<ButtonPress-1>", self._pan_start)
        self.canvas.bind("<B1-Motion>", self._pan_move)
        self.canvas.bind("<ButtonRelease-1>", self._pan_end)
        self._show_message("Selecciona las carpetas y carga un catálogo para comenzar.")

    @property
    def rotation(self) -> int:
        return self._rotation

    def close_image(self) -> None:
        if self._image is not None and self._image is not self._base:
            self._image.close()
        if self._base is not None:
            self._base.close()
        self._base = self._image = None
        self._photo = None

    def load(self, path: Path | None) -> None:
        self.close_image()
        self._path = path
        self._rotation = 0
        self._zoom = 1.0
        self._pan_x = self._pan_y = 0.0
        if path is None:
            self._show_message("No hay una fotografía seleccionada.")
            self.zoom_callback(0)
            return
        try:
            with Image.open(path) as opened:
                oriented = ImageOps.exif_transpose(opened)
                if oriented.mode not in ("RGB", "RGBA", "L"):
                    oriented = oriented.convert("RGB")
                self._base = oriented.copy()
            self._image = self._base
            self.canvas.after_idle(self.fit)
        except Exception as exc:
            self._show_message(f"No se pudo previsualizar esta imagen.\n\n{exc}\n\nPuedes omitirla.")
            self.zoom_callback(0)

    def _show_message(self, text: str) -> None:
        self.canvas.delete("all")
        width = max(self.canvas.winfo_width(), 500)
        height = max(self.canvas.winfo_height(), 350)
        self.canvas.create_text(
            width / 2, height / 2, text=text, fill="#E8ECEA", width=width - 80,
            justify="center", font=("Segoe UI", 11), tags="message",
        )

    def _viewport(self) -> tuple[int, int]:
        return max(1, self.canvas.winfo_width()), max(1, self.canvas.winfo_height())

    def _fit_scale(self) -> float:
        if self._image is None:
            return 1.0
        width, height = self._viewport()
        return min(width / self._image.width, height / self._image.height, self.MAX_ZOOM)

    def fit(self) -> None:
        if self._image is None:
            return
        self._zoom = max(self.MIN_ZOOM, self._fit_scale())
        self._pan_x = self._pan_y = 0.0
        self._render()

    def actual_size(self) -> None:
        width, height = self._viewport()
        self._set_zoom(1.0, width / 2, height / 2)

    def zoom_in(self) -> None:
        width, height = self._viewport()
        self._set_zoom(self._zoom * 1.25, width / 2, height / 2)

    def zoom_out(self) -> None:
        width, height = self._viewport()
        self._set_zoom(self._zoom / 1.25, width / 2, height / 2)

    def rotate_clockwise(self) -> None:
        if self._base is None:
            return
        self._rotation = (self._rotation + 90) % 360
        if self._image is not None and self._image is not self._base:
            self._image.close()
        transpose = {
            0: None, 90: Image.Transpose.ROTATE_270,
            180: Image.Transpose.ROTATE_180, 270: Image.Transpose.ROTATE_90,
        }[self._rotation]
        self._image = self._base if transpose is None else self._base.transpose(transpose)
        self.fit()

    def _origin(self, scale: float | None = None) -> tuple[float, float]:
        if self._image is None:
            return 0.0, 0.0
        scale = self._zoom if scale is None else scale
        view_width, view_height = self._viewport()
        scaled_width = self._image.width * scale
        scaled_height = self._image.height * scale
        x = (view_width - scaled_width) / 2 if scaled_width <= view_width else -self._pan_x
        y = (view_height - scaled_height) / 2 if scaled_height <= view_height else -self._pan_y
        return x, y

    def _clamp_pan(self) -> None:
        if self._image is None:
            return
        view_width, view_height = self._viewport()
        self._pan_x = min(max(0.0, self._pan_x), max(0.0, self._image.width * self._zoom - view_width))
        self._pan_y = min(max(0.0, self._pan_y), max(0.0, self._image.height * self._zoom - view_height))

    def _set_zoom(self, value: float, canvas_x: float, canvas_y: float) -> None:
        if self._image is None:
            return
        value = min(self.MAX_ZOOM, max(self.MIN_ZOOM, value))
        old_origin_x, old_origin_y = self._origin()
        image_x = min(max(0.0, (canvas_x - old_origin_x) / self._zoom), self._image.width)
        image_y = min(max(0.0, (canvas_y - old_origin_y) / self._zoom), self._image.height)
        self._zoom = value
        view_width, view_height = self._viewport()
        self._pan_x = image_x * value - canvas_x if self._image.width * value > view_width else 0.0
        self._pan_y = image_y * value - canvas_y if self._image.height * value > view_height else 0.0
        self._clamp_pan()
        self._render()

    def _render(self) -> None:
        if self._image is None:
            return
        self._clamp_pan()
        view_width, view_height = self._viewport()
        origin_x, origin_y = self._origin()
        scale = self._zoom
        left = max(0, math.floor((0 - origin_x) / scale))
        top = max(0, math.floor((0 - origin_y) / scale))
        right = min(self._image.width, math.ceil((view_width - origin_x) / scale))
        bottom = min(self._image.height, math.ceil((view_height - origin_y) / scale))
        if right <= left or bottom <= top:
            return
        crop = self._image.crop((left, top, right, bottom))
        render_width = max(1, round((right - left) * scale))
        render_height = max(1, round((bottom - top) * scale))
        if crop.size != (render_width, render_height):
            resampling = Image.Resampling.LANCZOS if scale < 1 else Image.Resampling.BICUBIC
            resized = crop.resize((render_width, render_height), resampling)
            crop.close()
            crop = resized
        self._photo = ImageTk.PhotoImage(crop)
        crop.close()
        self.canvas.delete("all")
        self.canvas.create_image(
            origin_x + left * scale, origin_y + top * scale,
            image=self._photo, anchor="nw",
        )
        self.zoom_callback(round(self._zoom * 100))

    def _on_wheel(self, event: tk.Event) -> str:
        if self._image is None:
            return "break"
        factor = 1.20 if event.delta > 0 else 1 / 1.20
        self._set_zoom(self._zoom * factor, event.x, event.y)
        return "break"

    def _on_linux_wheel(self, event: tk.Event, direction: int) -> str:
        if self._image is not None:
            factor = 1.20 if direction > 0 else 1 / 1.20
            self._set_zoom(self._zoom * factor, event.x, event.y)
        return "break"

    def _pan_start(self, event: tk.Event) -> None:
        self._drag_start = (event.x, event.y, self._pan_x, self._pan_y)

    def _pan_move(self, event: tk.Event) -> None:
        if self._drag_start is None or self._image is None:
            return
        start_x, start_y, pan_x, pan_y = self._drag_start
        self._pan_x = pan_x - (event.x - start_x)
        self._pan_y = pan_y - (event.y - start_y)
        self._render()

    def _pan_end(self, _event: tk.Event) -> None:
        self._drag_start = None

    def _on_configure(self, _event: tk.Event) -> None:
        if self._resize_job is not None:
            self.after_cancel(self._resize_job)
        self._resize_job = self.after(90, self._redraw_after_resize)

    def _redraw_after_resize(self) -> None:
        self._resize_job = None
        self._render()


class ThumbnailStrip(ttk.Frame):
    """Franja horizontal con carga progresiva y caché limitada."""

    CELL_WIDTH = 132
    CELL_HEIGHT = 112
    CACHE_LIMIT = 40

    def __init__(self, master: tk.Misc, on_select: Callable[[Path], None]):
        super().__init__(master, style="Panel.TFrame")
        self.on_select = on_select
        self.canvas = tk.Canvas(self, height=self.CELL_HEIGHT, background=WHITE, highlightthickness=0)
        self.scrollbar = ttk.Scrollbar(self, orient="horizontal", command=self._xview)
        self.canvas.configure(xscrollcommand=self.scrollbar.set)
        self.canvas.pack(fill="both", expand=True)
        self.scrollbar.pack(fill="x")
        self.items: list[Path] = []
        self.selected: Path | None = None
        self.status_getter: Callable[[Path], str] = lambda _path: "pending"
        self.cache: OrderedDict[str, Image.Image] = OrderedDict()
        self.tk_images: dict[str, ImageTk.PhotoImage] = {}
        self.queue: deque[Path] = deque()
        self.queued: set[str] = set()
        self.generation = 0
        self.loading = False
        self.canvas.bind("<Configure>", lambda _event: self._redraw())
        self.canvas.bind("<Button-1>", self._click)
        self.canvas.bind("<MouseWheel>", self._wheel)
        self.canvas.bind("<Shift-MouseWheel>", self._wheel)
        self.canvas.bind("<Button-4>", lambda event: self._linux_wheel(event, -1))
        self.canvas.bind("<Button-5>", lambda event: self._linux_wheel(event, 1))

    def set_items(
        self, items: list[Path], selected: Path | None,
        status_getter: Callable[[Path], str],
    ) -> None:
        self.items = list(items)
        self.selected = selected
        self.status_getter = status_getter
        self.generation += 1
        self.queue.clear()
        self.queued.clear()
        self.loading = False
        self.canvas.configure(scrollregion=(0, 0, len(items) * self.CELL_WIDTH, self.CELL_HEIGHT))
        self._redraw()

    def select(self, path: Path | None, reveal: bool = True) -> None:
        self.selected = path
        if reveal and path in self.items:
            index = self.items.index(path)
            total_width = max(1, len(self.items) * self.CELL_WIDTH)
            left = index * self.CELL_WIDTH
            right = left + self.CELL_WIDTH
            visible_left = self.canvas.canvasx(0)
            visible_right = visible_left + self.canvas.winfo_width()
            if left < visible_left:
                self.canvas.xview_moveto(left / total_width)
            elif right > visible_right:
                self.canvas.xview_moveto(max(0.0, (right - self.canvas.winfo_width()) / total_width))
        self._redraw()

    def _xview(self, *args: str) -> None:
        self.canvas.xview(*args)
        self._redraw()

    def _wheel(self, event: tk.Event) -> str:
        self.canvas.xview_scroll(-1 if event.delta > 0 else 1, "units")
        self._redraw()
        return "break"

    def _linux_wheel(self, _event: tk.Event, direction: int) -> str:
        self.canvas.xview_scroll(direction, "units")
        self._redraw()
        return "break"

    def _visible_range(self) -> tuple[int, int]:
        if not self.items:
            return 0, 0
        left = max(0, int(self.canvas.canvasx(0) // self.CELL_WIDTH) - 1)
        right = min(len(self.items), int((self.canvas.canvasx(self.canvas.winfo_width()) // self.CELL_WIDTH) + 2))
        return left, right

    def _redraw(self) -> None:
        self.canvas.delete("thumb")
        self.tk_images.clear()
        start, end = self._visible_range()
        for index in range(start, end):
            path = self.items[index]
            key = normalized(path)
            x = index * self.CELL_WIDTH + 5
            selected = self.selected == path
            status = self.status_getter(path)
            outline = GREEN if selected else "#D7DDD9"
            self.canvas.create_rectangle(
                x, 4, x + 122, 102, fill=GREEN_PALE if selected else WHITE,
                outline=outline, width=3 if selected else 1, tags="thumb",
            )
            image = self.cache.get(key)
            if image is not None:
                self.cache.move_to_end(key)
                photo = ImageTk.PhotoImage(image)
                self.tk_images[key] = photo
                self.canvas.create_image(x + 61, 42, image=photo, anchor="center", tags="thumb")
            else:
                self.canvas.create_text(x + 61, 42, text="Cargando…", fill=MUTED,
                                        font=("Segoe UI", 8), tags="thumb")
                if key not in self.queued:
                    self.queue.append(path)
                    self.queued.add(key)
            color = {"registered": GREEN, "skipped": GOLD, "pending": MUTED}[status]
            label = {"registered": "Registrada", "skipped": "Omitida", "pending": "Pendiente"}[status]
            self.canvas.create_text(x + 61, 80, text=label, fill=color,
                                    font=("Segoe UI", 8, "bold"), tags="thumb")
            short_name = path.name if len(path.name) <= 20 else path.name[:17] + "…"
            self.canvas.create_text(x + 61, 94, text=short_name, fill=INK,
                                    font=("Segoe UI", 7), tags="thumb")
        if self.queue and not self.loading:
            self.loading = True
            generation = self.generation
            self.after(15, lambda: self._load_one(generation))

    def _placeholder(self) -> Image.Image:
        image = Image.new("RGB", (112, 68), "#E6E9E7")
        draw = ImageDraw.Draw(image)
        draw.line((42, 22, 70, 50), fill=RED, width=3)
        draw.line((70, 22, 42, 50), fill=RED, width=3)
        return image

    def _load_one(self, generation: int) -> None:
        if generation != self.generation:
            self.loading = False
            return
        if not self.queue:
            self.loading = False
            return
        path = self.queue.popleft()
        key = normalized(path)
        self.queued.discard(key)
        try:
            with Image.open(path) as opened:
                image = ImageOps.exif_transpose(opened)
                if image.mode not in ("RGB", "RGBA"):
                    image = image.convert("RGB")
                image = ImageOps.contain(image, (112, 68), Image.Resampling.LANCZOS).copy()
        except Exception:
            image = self._placeholder()
        self.cache[key] = image
        self.cache.move_to_end(key)
        while len(self.cache) > self.CACHE_LIMIT:
            _, old = self.cache.popitem(last=False)
            old.close()
        self.loading = False
        self._redraw()

    def _click(self, event: tk.Event) -> None:
        index = int(self.canvas.canvasx(event.x) // self.CELL_WIDTH)
        if 0 <= index < len(self.items):
            self.on_select(self.items[index])


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Catálogo fotográfico de pozos")
        self.geometry("1280x820")
        self.minsize(780, 560)
        self.configure(background=SAND)
        self.protocol("WM_DELETE_WINDOW", self._close)
        self.catalog: Catalog | None = None
        self.current: Path | None = None
        self.saving = False
        self.editing_path: Path | None = None
        self.settings_open = True
        self.header_logo: ImageTk.PhotoImage | None = None
        self._sash_job: str | None = None

        self.source_var = tk.StringVar()
        self.output_var = tk.StringVar()
        self.site_var = tk.StringVar(value="H-96-AM")
        self.pit_var = tk.StringVar()
        self.level_var = tk.StringVar()
        self.view_var = tk.StringVar()
        self.keep_var = tk.BooleanVar(value=True)
        self.zoom_var = tk.StringVar(value="—")
        self.proposed_var = tk.StringVar(value="Completa Pozo y Nivel para ver el nombre previsto.")
        self.progress_var = tk.StringVar(value="Registradas: 0 · Pendientes: 0 · Omitidas por revisar: 0")
        self.status_var = tk.StringVar(value="Selecciona las carpetas de origen y salida.")

        self._configure_styles()
        self._build_header()
        self._build_settings()
        self._build_progress()
        self._build_workspace()
        self._build_thumbnails()
        self._build_statusbar()
        self._bind_shortcuts()
        self.pit_var.trace_add("write", lambda *_args: self._form_changed())
        self.level_var.trace_add("write", lambda *_args: self._form_changed())
        self.site_var.trace_add("write", lambda *_args: self._form_changed())
        self._update_actions()
        # Tk aplica la geometría real después del primer ciclo de dibujo.
        self.after(180, self._set_initial_sash)

    def _configure_styles(self) -> None:
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("Root.TFrame", background=SAND)
        style.configure("Panel.TFrame", background=WHITE)
        style.configure("TLabel", background=SAND, foreground=INK, font=("Segoe UI", 10))
        style.configure("Panel.TLabel", background=WHITE, foreground=INK, font=("Segoe UI", 10))
        style.configure("Muted.Panel.TLabel", background=WHITE, foreground=MUTED, font=("Segoe UI", 9))
        style.configure("Title.TLabel", background=SAND, foreground=INK, font=("Segoe UI Semibold", 17))
        style.configure("Subtitle.TLabel", background=SAND, foreground=MUTED, font=("Segoe UI", 9))
        style.configure("Section.Panel.TLabel", background=WHITE, foreground=GREEN, font=("Segoe UI Semibold", 12))
        style.configure("Required.Panel.TLabel", background=WHITE, foreground=INK, font=("Segoe UI Semibold", 10))
        style.configure("Accent.TButton", font=("Segoe UI Semibold", 10), padding=(15, 9), background=GREEN, foreground=WHITE)
        style.map("Accent.TButton", background=[("active", GREEN_DARK), ("disabled", "#A9B7B2")])
        style.configure("Secondary.TButton", font=("Segoe UI", 9), padding=(10, 7))
        style.configure("Tool.TButton", font=("Segoe UI", 9), padding=(8, 5))
        style.configure("Mandatory.TEntry", fieldbackground="#FBFCF8", bordercolor=GREEN,
                        lightcolor=GREEN, darkcolor=GREEN, padding=6)
        style.configure("Optional.TEntry", fieldbackground=WHITE, padding=6)
        style.configure("Green.Horizontal.TProgressbar", troughcolor="#DDE5E1", background=GREEN,
                        bordercolor=SAND, lightcolor=GREEN, darkcolor=GREEN)

    def _build_header(self) -> None:
        header = ttk.Frame(self, style="Root.TFrame", padding=(16, 10, 16, 8))
        header.grid(row=0, column=0, sticky="ew")
        header.columnconfigure(1, weight=1)
        if ASSET_PATH.is_file():
            try:
                with Image.open(ASSET_PATH) as opened:
                    logo = ImageOps.contain(opened.convert("RGBA"), (54, 54), Image.Resampling.LANCZOS)
                    self.header_logo = ImageTk.PhotoImage(logo)
                ttk.Label(header, image=self.header_logo, background=SAND).grid(
                    row=0, column=0, rowspan=2, padx=(0, 12))
            except Exception:
                pass
        ttk.Label(header, text="Catálogo fotográfico de pozos", style="Title.TLabel").grid(
            row=0, column=1, sticky="sw")
        ttk.Label(header,
                  text="Registro y preparación de fotografías para entrega al Museo Nacional",
                  style="Subtitle.TLabel").grid(row=1, column=1, sticky="nw")
        self.settings_button = ttk.Button(
            header, text="Ocultar configuración  ▴", style="Secondary.TButton",
            command=self._toggle_settings)
        self.settings_button.grid(row=0, column=2, rowspan=2, sticky="e")

    def _build_settings(self) -> None:
        self.settings_frame = ttk.Frame(self, style="Panel.TFrame", padding=(16, 10))
        self.settings_frame.grid(row=1, column=0, sticky="ew", padx=16, pady=(0, 8))
        self.settings_frame.columnconfigure(1, weight=1)
        rows = (("Fotos descargadas", self.source_var, self._choose_source),
                ("Carpeta de salida", self.output_var, self._choose_output))
        for row_index, (label, variable, command) in enumerate(rows):
            ttk.Label(self.settings_frame, text=label, style="Panel.TLabel").grid(
                row=row_index, column=0, sticky="w", padx=(0, 10), pady=4)
            ttk.Entry(self.settings_frame, textvariable=variable, style="Optional.TEntry").grid(
                row=row_index, column=1, sticky="ew", pady=4)
            ttk.Button(self.settings_frame, text="Elegir…", command=command,
                       style="Secondary.TButton").grid(row=row_index, column=2, padx=(8, 0), pady=4)
        ttk.Label(self.settings_frame, text="Código del sitio", style="Panel.TLabel").grid(
            row=2, column=0, sticky="w", pady=4)
        ttk.Entry(self.settings_frame, textvariable=self.site_var, width=20,
                  style="Mandatory.TEntry").grid(row=2, column=1, sticky="w", pady=4)
        ttk.Button(self.settings_frame, text="Cargar / reanudar", command=self._load_catalog,
                   style="Accent.TButton").grid(row=2, column=2, padx=(8, 0), pady=4)

    def _build_progress(self) -> None:
        frame = ttk.Frame(self, style="Root.TFrame", padding=(16, 0, 16, 8))
        frame.grid(row=2, column=0, sticky="ew")
        frame.columnconfigure(0, weight=1)
        self.progress = ttk.Progressbar(frame, style="Green.Horizontal.TProgressbar", maximum=1, value=0)
        self.progress.grid(row=0, column=0, sticky="ew", padx=(0, 12))
        ttk.Label(frame, textvariable=self.progress_var, style="Subtitle.TLabel").grid(
            row=0, column=1, sticky="e")

    def _build_workspace(self) -> None:
        self.paned = ttk.Panedwindow(self, orient="horizontal")
        self.paned.grid(row=3, column=0, sticky="nsew", padx=16)
        self.paned.bind("<Configure>", self._paned_configure)
        self.rowconfigure(3, weight=1)
        self.columnconfigure(0, weight=1)
        left = ttk.Frame(self.paned, style="Panel.TFrame", padding=8)
        right_holder = ttk.Frame(self.paned, style="Panel.TFrame")
        self.paned.add(left, weight=4)
        self.paned.add(right_holder, weight=2)
        toolbar = ttk.Frame(left, style="Panel.TFrame")
        toolbar.pack(fill="x", pady=(0, 7))
        ttk.Button(toolbar, text="−", command=self._zoom_out, style="Tool.TButton", width=3).pack(side="left")
        ttk.Button(toolbar, text="+", command=self._zoom_in, style="Tool.TButton", width=3).pack(side="left", padx=(4, 0))
        ttk.Button(toolbar, text="Ajustar", command=self._fit, style="Tool.TButton").pack(side="left", padx=(8, 0))
        ttk.Button(toolbar, text="100 %", command=self._actual, style="Tool.TButton").pack(side="left", padx=(4, 0))
        ttk.Button(toolbar, text="Girar 90°", command=self._rotate, style="Tool.TButton").pack(side="left", padx=(8, 0))
        ttk.Label(toolbar, textvariable=self.zoom_var, style="Panel.TLabel",
                  font=("Segoe UI Semibold", 10)).pack(side="right", padx=6)
        self.viewer = ImageViewer(left, lambda percent: self.zoom_var.set(f"{percent} %" if percent else "—"))
        self.viewer.pack(fill="both", expand=True)

        form_area = ttk.Frame(right_holder, style="Panel.TFrame")
        form_area.pack(fill="both", expand=True)
        self.form_canvas = tk.Canvas(form_area, background=WHITE, highlightthickness=0, width=350)
        form_scroll = ttk.Scrollbar(form_area, orient="vertical", command=self.form_canvas.yview)
        self.form_canvas.configure(yscrollcommand=form_scroll.set)
        form_scroll.pack(side="right", fill="y")
        self.form_canvas.pack(side="left", fill="both", expand=True)
        self.form = ttk.Frame(self.form_canvas, style="Panel.TFrame", padding=16)
        self.form_window = self.form_canvas.create_window((0, 0), window=self.form, anchor="nw")
        self.form.columnconfigure(0, weight=1)
        self.form.bind("<Configure>", self._form_configure)
        self.form_canvas.bind("<Configure>", self._form_canvas_configure)
        self.form_canvas.bind("<Enter>", self._form_enter)
        self.form_canvas.bind("<Leave>", self._form_leave)
        self._build_form_fields()
        self._build_form_actions(right_holder)

    def _build_form_fields(self) -> None:
        ttk.Label(self.form, text="Datos de registro", style="Section.Panel.TLabel").grid(
            row=0, column=0, sticky="w")
        self.current_label = ttk.Label(self.form, text="Sin fotografía seleccionada",
                                       style="Muted.Panel.TLabel", wraplength=320)
        self.current_label.grid(row=1, column=0, sticky="ew", pady=(3, 14))
        ttk.Label(self.form, text="Pozo  ·  obligatorio", style="Required.Panel.TLabel").grid(
            row=2, column=0, sticky="w")
        self.pit_entry = ttk.Entry(self.form, textvariable=self.pit_var,
                                   style="Mandatory.TEntry", font=("Segoe UI", 12))
        self.pit_entry.grid(row=3, column=0, sticky="ew", pady=(3, 11))
        ttk.Label(self.form, text="Nivel  ·  obligatorio", style="Required.Panel.TLabel").grid(
            row=4, column=0, sticky="w")
        self.level_entry = ttk.Entry(self.form, textvariable=self.level_var,
                                     style="Mandatory.TEntry", font=("Segoe UI", 12))
        self.level_entry.grid(row=5, column=0, sticky="ew", pady=(3, 13))
        ttk.Label(self.form, text="Nombre previsto", style="Panel.TLabel").grid(
            row=6, column=0, sticky="w")
        self.proposed_label = ttk.Label(
            self.form, textvariable=self.proposed_var, style="Muted.Panel.TLabel",
            wraplength=320, justify="left")
        self.proposed_label.grid(row=7, column=0, sticky="ew", pady=(3, 14))
        ttk.Separator(self.form).grid(row=8, column=0, sticky="ew", pady=(0, 13))
        ttk.Label(self.form, text="Vista / detalle  ·  opcional", style="Panel.TLabel").grid(
            row=9, column=0, sticky="w")
        ttk.Entry(self.form, textvariable=self.view_var, style="Optional.TEntry").grid(
            row=10, column=0, sticky="ew", pady=(3, 11))
        ttk.Label(self.form, text="Observaciones  ·  opcional", style="Panel.TLabel").grid(
            row=11, column=0, sticky="w")
        self.notes = tk.Text(
            self.form, height=5, wrap="word", relief="solid", borderwidth=1,
            highlightthickness=1, highlightbackground="#CDD5D1", highlightcolor=GREEN,
            font=("Segoe UI", 10), foreground=INK, background=WHITE)
        self.notes.grid(row=12, column=0, sticky="ew", pady=(3, 11))

    def _build_form_actions(self, parent: tk.Misc) -> None:
        actions = ttk.Frame(parent, style="Panel.TFrame", padding=(16, 8, 16, 12))
        actions.pack(fill="x", side="bottom")
        actions.columnconfigure(0, weight=1)
        ttk.Separator(actions).grid(row=0, column=0, sticky="ew", pady=(0, 8))
        ttk.Checkbutton(actions, text="Mantener pozo y nivel para la siguiente foto",
                        variable=self.keep_var).grid(row=1, column=0, sticky="w", pady=(0, 8))
        self.save_button = ttk.Button(actions, text="Guardar y siguiente",
                                      command=self._save, style="Accent.TButton")
        self.save_button.grid(row=2, column=0, sticky="ew")
        self.skip_button = ttk.Button(actions, text="Omitir por ahora",
                                      command=self._toggle_skip, style="Secondary.TButton")
        self.skip_button.grid(row=3, column=0, sticky="ew", pady=(6, 0))
        self.help_label = ttk.Label(
            actions,
            text="* La copia conserva su formato. El zoom y el giro no modifican el archivo.",
            style="Muted.Panel.TLabel", wraplength=320, justify="left")
        self.help_label.grid(row=4, column=0, sticky="ew", pady=(8, 0))

    def _build_thumbnails(self) -> None:
        frame = ttk.Frame(self, style="Panel.TFrame", padding=(8, 6))
        frame.grid(row=4, column=0, sticky="ew", padx=16, pady=(8, 0))
        ttk.Label(frame, text="Fotografías", style="Section.Panel.TLabel").pack(
            anchor="w", padx=3, pady=(0, 3))
        self.thumbnails = ThumbnailStrip(frame, self._select_photo)
        self.thumbnails.pack(fill="x")

    def _build_statusbar(self) -> None:
        self.status_label = tk.Label(
            self, textvariable=self.status_var, background=GREEN_PALE, foreground=INK,
            anchor="w", padx=16, pady=7, font=("Segoe UI", 9))
        self.status_label.grid(row=5, column=0, sticky="ew", pady=(8, 0))

    def _bind_shortcuts(self) -> None:
        self.bind_all("<Control-plus>", self._shortcut_zoom_in)
        self.bind_all("<Control-KP_Add>", self._shortcut_zoom_in)
        self.bind_all("<Control-equal>", self._shortcut_zoom_in)
        self.bind_all("<Control-minus>", self._shortcut_zoom_out)
        self.bind_all("<Control-KP_Subtract>", self._shortcut_zoom_out)
        self.bind_all("<Control-Key-0>", self._shortcut_fit)
        self.bind_all("<Control-Key-1>", self._shortcut_actual)
        # En macOS, Command es el modificador habitual de la aplicación.
        self.bind_all("<Command-plus>", self._shortcut_zoom_in)
        self.bind_all("<Command-equal>", self._shortcut_zoom_in)
        self.bind_all("<Command-minus>", self._shortcut_zoom_out)
        self.bind_all("<Command-Key-0>", self._shortcut_fit)
        self.bind_all("<Command-Key-1>", self._shortcut_actual)

    def _shortcut_allowed(self) -> bool:
        focus = self.focus_get()
        return not isinstance(focus, (tk.Entry, tk.Text, ttk.Entry, ttk.Combobox, ttk.Spinbox))

    def _shortcut_zoom_in(self, _event: tk.Event) -> str | None:
        if self._shortcut_allowed():
            self.viewer.zoom_in()
            return "break"
        return None

    def _shortcut_zoom_out(self, _event: tk.Event) -> str | None:
        if self._shortcut_allowed():
            self.viewer.zoom_out()
            return "break"
        return None

    def _shortcut_fit(self, _event: tk.Event) -> str | None:
        if self._shortcut_allowed():
            self.viewer.fit()
            return "break"
        return None

    def _shortcut_actual(self, _event: tk.Event) -> str | None:
        if self._shortcut_allowed():
            self.viewer.actual_size()
            return "break"
        return None

    def _set_initial_sash(self) -> None:
        try:
            self.paned.sashpos(0, max(500, int(self.winfo_width() * 0.68)))
        except tk.TclError:
            pass

    def _paned_configure(self, event: tk.Event) -> None:
        """Mantiene ambos paneles utilizables, incluso en la ventana mínima."""
        if self._sash_job is not None:
            self.after_cancel(self._sash_job)
        self._sash_job = self.after(20, self._enforce_pane_bounds)

    def _enforce_pane_bounds(self) -> None:
        self._sash_job = None
        try:
            width = self.paned.winfo_width()
            if width < 650:
                return
            position = self.paned.sashpos(0)
            minimum_left = 340
            minimum_right = 285
            bounded = min(max(position, minimum_left), width - minimum_right)
            if bounded != position:
                self.paned.sashpos(0, bounded)
        except tk.TclError:
            pass

    def _toggle_settings(self) -> None:
        self.settings_open = not self.settings_open
        if self.settings_open:
            self.settings_frame.grid()
            self.settings_button.configure(text="Ocultar configuración  ▴")
        else:
            self.settings_frame.grid_remove()
            self.settings_button.configure(text="Mostrar configuración  ▾")

    def _form_configure(self, _event: tk.Event) -> None:
        self.form_canvas.configure(scrollregion=self.form_canvas.bbox("all"))

    def _form_canvas_configure(self, event: tk.Event) -> None:
        self.form_canvas.itemconfigure(self.form_window, width=event.width)
        wrap = max(140, event.width - 38)
        self.current_label.configure(wraplength=wrap)
        self.proposed_label.configure(wraplength=wrap)
        self.help_label.configure(wraplength=wrap)

    def _form_enter(self, _event: tk.Event) -> None:
        self.form_canvas.bind_all("<MouseWheel>", self._form_wheel)
        self.form_canvas.bind_all("<Button-4>", lambda event: self._form_linux_wheel(event, -1))
        self.form_canvas.bind_all("<Button-5>", lambda event: self._form_linux_wheel(event, 1))

    def _form_leave(self, _event: tk.Event) -> None:
        self.form_canvas.unbind_all("<MouseWheel>")
        self.form_canvas.unbind_all("<Button-4>")
        self.form_canvas.unbind_all("<Button-5>")
        self.viewer.canvas.bind("<MouseWheel>", self.viewer._on_wheel)

    def _form_wheel(self, event: tk.Event) -> str:
        self.form_canvas.yview_scroll(-1 if event.delta > 0 else 1, "units")
        return "break"

    def _form_linux_wheel(self, _event: tk.Event, direction: int) -> str:
        self.form_canvas.yview_scroll(direction, "units")
        return "break"

    def _choose_source(self) -> None:
        chosen = filedialog.askdirectory(title="Carpeta con fotografías descargadas")
        if chosen:
            self.source_var.set(chosen)

    def _choose_output(self) -> None:
        chosen = filedialog.askdirectory(title="Carpeta de trabajo y salida")
        if chosen:
            self.output_var.set(chosen)

    def _load_catalog(self) -> None:
        try:
            source = Path(self.source_var.get().strip())
            output = Path(self.output_var.get().strip())
            if not source.is_dir() or not output.is_dir():
                raise ValueError("Selecciona dos carpetas existentes.")
            if source.resolve() == output.resolve():
                raise ValueError("El origen y la salida deben ser carpetas distintas.")
            site = Catalog.validate_identifier(self.site_var.get(), "Código del sitio")
            catalog = Catalog(source, output, site)
            if catalog.rows and catalog.rows[0]["Sitio"] != site:
                raise ValueError("El catálogo existente usa otro código de sitio.")
            self.catalog = catalog
            self.thumbnails.set_items(catalog.files, None, catalog.status_of)
            selected = self._resume_path()
            self._select_photo(selected, remember=False)
            self._update_progress()
            self._set_status(f"Catálogo listo: {len(catalog.files)} fotografía(s) compatible(s).", "success")
            if self.settings_open:
                self._toggle_settings()
        except Exception as exc:
            messagebox.showerror("No se pudo cargar", str(exc), parent=self)
            self._set_status("No se pudo cargar el catálogo. Revisa la configuración.", "error")

    def _resume_path(self) -> Path | None:
        assert self.catalog is not None
        if self.catalog.last_path:
            for path in self.catalog.files:
                if normalized(path) == self.catalog.last_path and self.catalog.status_of(path) != "registered":
                    return path
        remaining = self.catalog.remaining()
        if remaining:
            return remaining[0]
        omitted = self.catalog.omitted()
        if omitted:
            return omitted[0]
        return self.catalog.files[-1] if self.catalog.files else None

    def _select_photo(self, path: Path | None, remember: bool = True) -> None:
        if path is not None and self.catalog is not None and path not in self.catalog.files:
            return
        if self.editing_path is not None and path != self.editing_path:
            if not messagebox.askyesno(
                "Cancelar corrección",
                "Hay una corrección sin guardar. ¿Quieres descartarla y cambiar de fotografía?",
                parent=self,
            ):
                return
            self._cancel_correction(clear_status=False)
        self.current = path
        self.viewer.load(path)
        self.thumbnails.select(path)
        if path is None:
            self.current_label.configure(text="No hay fotografías compatibles.")
        elif self.catalog is not None:
            status = self.catalog.status_of(path)
            status_text = {"registered": "Registrada", "skipped": "Omitida por revisar",
                           "pending": "Pendiente"}[status]
            self.current_label.configure(text=f"{path.name}\n{status_text}")
            if remember:
                try:
                    self.catalog.remember_current(path)
                except OSError:
                    pass
        self._form_changed()

    def _form_changed(self) -> None:
        if self.catalog is None or self.current is None:
            self.proposed_var.set("Completa Pozo y Nivel para ver el nombre previsto.")
            self._update_actions()
            return
        status = self.catalog.status_of(self.current)
        if self.editing_path == self.current:
            try:
                self.proposed_var.set(self.catalog.proposed_correction_name(
                    self.current, self.pit_var.get(), self.level_var.get()
                ))
            except ValueError:
                self.proposed_var.set("Completa Pozo y Nivel con letras, números o guiones.")
        elif status == "registered":
            self.proposed_var.set("Esta fotografía ya fue registrada.")
        elif status == "skipped":
            self.proposed_var.set("Recupérala a pendientes para poder registrarla.")
        else:
            try:
                self.proposed_var.set(
                    self.catalog.proposed_name(self.current, self.pit_var.get(), self.level_var.get()))
            except ValueError:
                self.proposed_var.set("Completa Pozo y Nivel con letras, números o guiones.")
        self._update_actions()

    def _valid_form(self) -> bool:
        if self.catalog is None or self.current is None:
            return False
        status = self.catalog.status_of(self.current)
        if self.editing_path == self.current:
            if status != "registered":
                return False
        elif status != "pending":
            return False
        try:
            Catalog.validate_identifier(self.pit_var.get(), "Pozo")
            Catalog.validate_identifier(self.level_var.get(), "Nivel")
            return True
        except ValueError:
            return False

    def _update_actions(self) -> None:
        self.save_button.configure(
            text="Guardar corrección" if self.editing_path is not None else "Guardar y siguiente"
        )
        self.save_button.configure(state="normal" if self._valid_form() and not self.saving else "disabled")
        if self.catalog is None or self.current is None:
            self.skip_button.configure(text="Omitir por ahora", state="disabled")
            return
        status = self.catalog.status_of(self.current)
        if self.editing_path == self.current:
            self.skip_button.configure(
                text="Cancelar corrección", state="normal" if not self.saving else "disabled"
            )
        elif status == "skipped":
            self.skip_button.configure(text="Recuperar a pendientes",
                                       state="normal" if not self.saving else "disabled")
        elif status == "pending":
            self.skip_button.configure(text="Omitir por ahora",
                                       state="normal" if not self.saving else "disabled")
        else:
            self.skip_button.configure(
                text="Corregir datos", state="normal" if not self.saving else "disabled"
            )

    def _next_pending_after(self, path: Path | None) -> Path | None:
        assert self.catalog is not None
        remaining = self.catalog.remaining()
        if not remaining:
            return None
        if path not in self.catalog.files:
            return remaining[0]
        start = self.catalog.files.index(path)
        for offset in range(1, len(self.catalog.files) + 1):
            candidate = self.catalog.files[(start + offset) % len(self.catalog.files)]
            if self.catalog.status_of(candidate) == "pending":
                return candidate
        return remaining[0]

    def _save(self) -> None:
        if self.catalog is None or self.current is None or not self._valid_form():
            return
        if self.editing_path == self.current:
            self._save_correction()
            return
        saved_path = self.current
        self.saving = True
        self._update_actions()
        self._set_status("Guardando copia, Excel y estado…", "info")
        self.update_idletasks()
        committed_warning: str | None = None
        predicted_name = self.catalog.proposed_name(
            saved_path, self.pit_var.get(), self.level_var.get())
        try:
            row = self.catalog.save(
                saved_path, self.pit_var.get(), self.level_var.get(),
                self.view_var.get(), self.notes.get("1.0", "end-1c"))
        except SaveCommittedWarning as exc:
            committed_warning = str(exc)
            row = {"Nombre_archivo": predicted_name}
        except Exception as exc:
            self.saving = False
            self._update_actions()
            self._set_status(
                f"No se guardó nada: {exc}. Los datos del formulario se conservaron.",
                "error",
            )
            messagebox.showerror(
                "No se pudo guardar",
                f"{exc}\n\nSi el Excel está abierto, ciérralo y vuelve a intentar. "
                "Pozo, nivel y observaciones permanecen en pantalla.", parent=self)
            return

        self.saving = False
        if not self.keep_var.get():
            self.pit_var.set("")
            self.level_var.set("")
        self.view_var.set("")
        self.notes.delete("1.0", "end")
        self.thumbnails.set_items(self.catalog.files, saved_path, self.catalog.status_of)
        next_path = self._next_pending_after(saved_path)
        self._select_photo(next_path if next_path is not None else saved_path)
        self._update_progress()
        if committed_warning:
            self._set_status(committed_warning, "warning")
            messagebox.showwarning("Guardado con aviso", committed_warning, parent=self)
        else:
            self._set_status(f"Guardado correctamente: {row['Nombre_archivo']}", "success")
        self.pit_entry.focus_set()

    def _begin_correction(self) -> None:
        if self.catalog is None or self.current is None:
            return
        located = self.catalog.row_for_path(self.current)
        if located is None:
            messagebox.showerror(
                "No se puede corregir",
                "No se pudo asociar esta fotografía con una única fila del Excel. "
                "Esto puede ocurrir en un catálogo antiguo con nombres originales repetidos.",
                parent=self,
            )
            return
        _, row = located
        self.editing_path = self.current
        self.pit_var.set(row["Pozo"])
        self.level_var.set(row["Nivel"])
        self.view_var.set(row["Vista_o_detalle"])
        self.notes.delete("1.0", "end")
        self.notes.insert("1.0", row["Observaciones"])
        self.current_label.configure(text=f"{self.current.name}\nCorrigiendo registro")
        self._set_status(
            "Modo corrección: revisa los datos y pulsa Guardar corrección.", "warning"
        )
        self._form_changed()
        self.pit_entry.focus_set()

    def _cancel_correction(self, clear_status: bool = True) -> None:
        self.editing_path = None
        self.pit_var.set("")
        self.level_var.set("")
        self.view_var.set("")
        self.notes.delete("1.0", "end")
        if self.current is not None and self.catalog is not None:
            self.current_label.configure(text=f"{self.current.name}\nRegistrada")
        if clear_status:
            self._set_status("Corrección cancelada; el registro no cambió.", "info")
        self._form_changed()

    def _save_correction(self) -> None:
        assert self.catalog is not None and self.current is not None
        path = self.current
        self.saving = True
        self._update_actions()
        self._set_status("Aplicando corrección en la copia, el Excel y el estado…", "info")
        self.update_idletasks()
        warning: str | None = None
        try:
            row = self.catalog.correct(
                path, self.pit_var.get(), self.level_var.get(),
                self.view_var.get(), self.notes.get("1.0", "end-1c"),
            )
        except CorrectionCommittedWarning as exc:
            warning = str(exc)
            located = self.catalog.row_for_path(path)
            row = located[1] if located else {"Nombre_archivo": "registro corregido"}
        except Exception as exc:
            self.saving = False
            self._update_actions()
            self._set_status(
                f"No se aplicó la corrección: {exc}. El formulario se conservó.",
                "error",
            )
            messagebox.showerror(
                "No se pudo corregir",
                f"{exc}\n\nSi el Excel está abierto, ciérralo y vuelve a intentar.",
                parent=self,
            )
            return

        self.saving = False
        self.editing_path = None
        # Una corrección no avanza a otra foto: no se arrastran sus valores al
        # formulario normal aunque la opción de repetición esté activada.
        self.pit_var.set("")
        self.level_var.set("")
        self.view_var.set("")
        self.notes.delete("1.0", "end")
        self.thumbnails.set_items(self.catalog.files, path, self.catalog.status_of)
        self._select_photo(path)
        self._update_progress()
        if warning:
            self._set_status(warning, "warning")
            messagebox.showwarning("Corrección guardada con aviso", warning, parent=self)
        else:
            self._set_status(f"Corrección guardada: {row['Nombre_archivo']}", "success")

    def _toggle_skip(self) -> None:
        if self.catalog is None or self.current is None or self.saving:
            return
        if self.editing_path == self.current:
            self._cancel_correction()
            return
        status = self.catalog.status_of(self.current)
        if status == "registered":
            self._begin_correction()
            return
        path = self.current
        try:
            if status == "skipped":
                self.catalog.set_skipped(path, False)
                self._set_status(f"Recuperada a pendientes: {path.name}", "success")
                target = path
            else:
                self.catalog.set_skipped(path, True)
                self._set_status(f"Omitida por ahora: {path.name}", "warning")
                target = self._next_pending_after(path) or path
            self.thumbnails.set_items(self.catalog.files, target, self.catalog.status_of)
            self._select_photo(target)
            self._update_progress()
        except Exception as exc:
            messagebox.showerror("No se pudo actualizar", str(exc), parent=self)
            self._set_status("No se cambió el estado de la fotografía.", "error")

    def _update_progress(self) -> None:
        if self.catalog is None:
            return
        total = len(self.catalog.files)
        registered = sum(1 for path in self.catalog.files
                         if self.catalog.status_of(path) == "registered")
        skipped = sum(1 for path in self.catalog.files
                      if self.catalog.status_of(path) == "skipped")
        pending = total - registered - skipped
        self.progress.configure(maximum=max(1, total), value=registered)
        self.progress_var.set(
            f"Registradas: {registered} · Pendientes: {pending} · Omitidas por revisar: {skipped}")
        self.thumbnails._redraw()

    def _set_status(self, text: str, kind: str = "info") -> None:
        colors = {
            "info": (GREEN_PALE, INK), "success": ("#DDECE5", GREEN_DARK),
            "warning": ("#F5E9D5", "#70501F"), "error": ("#F3DEDC", "#7C302C"),
        }
        background, foreground = colors[kind]
        self.status_var.set(text)
        self.status_label.configure(background=background, foreground=foreground)

    def _zoom_in(self) -> None:
        self.viewer.zoom_in()

    def _zoom_out(self) -> None:
        self.viewer.zoom_out()

    def _fit(self) -> None:
        self.viewer.fit()

    def _actual(self) -> None:
        self.viewer.actual_size()

    def _rotate(self) -> None:
        self.viewer.rotate_clockwise()

    def _close(self) -> None:
        if self.catalog is not None:
            try:
                self.catalog.remember_current(self.current)
            except OSError:
                pass
        self.viewer.close_image()
        for image in self.thumbnails.cache.values():
            image.close()
        self.destroy()


def main() -> None:
    enable_windows_dpi_awareness()
    App().mainloop()


if __name__ == "__main__":
    main()
